import sys

import openpyxl
import openpyxl.utils

def iCare_parse_columns(file_name):
    book = openpyxl.load_workbook(filename = file_name, read_only = True)

    sheet = book.active

    columns = []
    row = 3

    for i in range(1, sheet.max_column + 1):
        cell = "{}{}".format(openpyxl.utils.get_column_letter(i), row)
        columns.append(sheet[cell].value)

    return columns

def parse_row(sheet, column, row):
    vals = []
    row = str(row)
    for col in column:
        if (not sheet[col + row].value):
            vals.append("''")
        else:
            vals.append(sheet[col + row].value)
    return vals

def parse_xlsx(file_name, column_names):
    book = openpyxl.load_workbook(filename = file_name, read_only = True)
    sheet = book.active

    values = [parse_row(sheet, column_names, i)
                      for i in range(4, sheet.max_row + 1)]
    return values
